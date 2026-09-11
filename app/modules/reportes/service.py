"""Reglas y composición del módulo de reportes de eventos."""

from collections import defaultdict
from datetime import date
from io import BytesIO
from math import ceil
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.modules.auditoria.repository import AuditoriaRepository
from app.modules.eventos.models import EventoModalidad
from app.modules.maestros.models import TipoCalculoBeneficio
from app.modules.reportes.dto import (
    AcreditacionReporteResponse,
    AreaReporte,
    BeneficioReporteItem,
    BeneficioReporteResponse,
    CalidadDato,
    CategoriaFiltroOpcion,
    ClasificacionReporte,
    CoberturaResponse,
    DashboardReporteResponse,
    DetalleReporteItem,
    DetalleReporteListResponse,
    DiaProgramacionReporte,
    DisponibilidadCampo,
    DisponibilidadDatos,
    DistribucionCategoriaItem,
    DistribucionCategoriaResponse,
    DistribucionGrupoItem,
    DistribucionGrupoResponse,
    EmpresaReporteItem,
    EmpresaReporteListResponse,
    EventoReporteCabecera,
    EventoSelectorItem,
    EventoSelectorResponse,
    FiltroOpcion,
    LugarReporte,
    MetricaValor,
    ParticipanteReporteItem,
    ParticipanteReporteListResponse,
    ProgramacionReporteItem,
    ProgramacionReporteResponse,
    ReporteClasificacion,
    ReporteContexto,
    ReporteEventoFiltros,
    ReporteFiltrosResponse,
    ReporteKpis,
    ReporteKpisResponse,
    ReporteOrden,
    ReporteTipoParticipante,
    ResponsableProgramacionReporte,
    SerieDato,
    SerieMensualResponse,
    SerieReporte,
)
from app.modules.reportes.repository import EventoReporteDetalle, ReporteRepository
from app.modules.usuarios.models import Usuario
from app.modules.usuarios.repository import UsuarioRepository


MODULO_REPORTES = "REPORTES"
PERMISO_PII = "CONSULTAR_DATOS_PERSONALES_REPORTE"


class ReporteServiceError(Exception):
    codigo = "REPORTE_ERROR"
    campo: str | None = None


class EventoReporteNotFoundError(ReporteServiceError):
    codigo = "REPORTE_EVENTO_NO_ENCONTRADO"


class ReporteFiltroInconsistenteError(ReporteServiceError):
    codigo = "REPORTE_FILTRO_INCONSISTENTE"

    def __init__(self, message: str, *, campo: str) -> None:
        super().__init__(message)
        self.campo = campo


class ReporteExportLimitError(ReporteServiceError):
    codigo = "REPORTE_EXPORTACION_EXCEDE_LIMITE"


class ReporteService:
    SORTS: dict[str, set[str]] = {
        "eventos": {"nombre_evento", "estado", "fecha_inicio"},
        "empresas": {"nombre_empresa", "ruc", "participantes"},
        "participantes": {
            "apellidos",
            "nombres",
            "id_evento_contacto",
            "id_programacion_evento",
            "id_empresa",
        },
        "detalle": {
            "apellidos",
            "nombres",
            "id_evento_contacto",
            "id_programacion_evento",
            "id_empresa",
        },
    }

    METRICAS: dict[str, tuple[str, str]] = {
        "total_programaciones": ("programación", "Programaciones del evento."),
        "total_dias": ("día programado", "Días configurados en las programaciones."),
        "total_empresas_afiliadas": ("afiliación", "Afiliaciones empresa-evento."),
        "total_empresas_unicas": ("empresa", "Empresas afiliadas distintas."),
        "total_empresas_representadas": ("empresa", "Empresas con al menos una asistencia."),
        "total_participaciones": ("participación", "Registros EventoContacto."),
        "total_contactos_maestros": ("participación", "Participaciones vinculadas a Contacto."),
        "total_invitados": ("participación", "Invitados almacenados como snapshot."),
        "contactos_unicos": ("contacto", "Contactos maestros distintos; no deduplica invitados."),
        "total_asistencias": ("participación", "Participaciones con asistencia registrada."),
        "total_sin_asistencia": ("participación", "Participaciones sin asistencia registrada."),
        "asignaciones_beneficio": ("asignación", "Asignaciones de beneficio existentes."),
        "requieren_coordinacion": ("participación", "Participaciones que requieren coordinación."),
        "qr_generados": ("participación", "Participaciones con QR generado."),
        "qr_enviados": ("participación", "QR con fecha de envío."),
        "qr_pendientes": ("participación", "QR activo generado y aún no enviado."),
        "credenciales_impresas": ("participación", "Credenciales marcadas como impresas."),
        "credenciales_no_impresas": ("participación", "Credenciales pendientes de impresión."),
    }

    SERIES_FUENTE = {
        "dias_programados": "FECHA_DIA_PROGRAMACION",
        "programaciones_programadas": "PRIMER_DIA_PROGRAMACION",
        "empresas_programadas": "PRIMER_DIA_PROGRAMACION",
        "participaciones_programadas": "PRIMER_DIA_PROGRAMACION",
        "asistencias_programadas": "PRIMER_DIA_PROGRAMACION",
        "empresas_representadas_programadas": "PRIMER_DIA_PROGRAMACION",
        "asignaciones_beneficio": "FECHA_ASIGNACION",
        "qr_generados": "FECHA_GENERACION_QR",
        "qr_enviados": "FECHA_ENVIO_QR",
    }

    def __init__(self, db: AsyncSession) -> None:
        self.db = db
        self.reportes = ReporteRepository(db)
        self.usuarios = UsuarioRepository(db)
        self.auditoria = AuditoriaRepository(db)

    @staticmethod
    def _pages(total: int, page_size: int) -> int:
        return ceil(total / page_size) if total else 0

    @staticmethod
    def _coverage(kind: str, numerator: int, denominator: int) -> CoberturaResponse:
        return CoberturaResponse(
            tipo_universo=kind,
            numerador=numerator,
            denominador=denominator,
            porcentaje=round(numerator * 100 / denominator, 2) if denominator else None,
        )

    @classmethod
    def _metric(cls, name: str, value: int) -> MetricaValor:
        grain, definition = cls.METRICAS[name]
        return MetricaValor(valor=value, grano=grain, definicion=definition)

    @staticmethod
    def _availability() -> DisponibilidadDatos:
        return DisponibilidadDatos(
            confirmacion=DisponibilidadCampo(
                disponible=False,
                motivo="El sistema no persiste confirmación de asistencia.",
            ),
            asistencia_por_dia=DisponibilidadCampo(
                disponible=False,
                motivo="La asistencia se registra por participante y programación.",
            ),
            canal_asistencia_hibrida=DisponibilidadCampo(
                disponible=False,
                motivo="En modalidad HIBRIDO no se registra el canal individual.",
            ),
            excedentes=DisponibilidadCampo(
                disponible=False,
                motivo="No existe un registro estructurado de excedentes.",
            ),
        )

    @staticmethod
    def _warnings(filtros: ReporteEventoFiltros) -> list[str]:
        warnings: list[str] = []
        if filtros.alcance.value == "HISTORICO":
            warnings.append(
                "El alcance histórico incluye filas inactivas existentes, pero no recupera eliminaciones físicas."
            )
        if filtros.clasificacion == ReporteClasificacion.FECHA_PROGRAMACION:
            warnings.append(
                "La clasificación se resuelve usando el primer día activo de cada programación."
            )
        return warnings

    @staticmethod
    def _context(filtros: ReporteEventoFiltros) -> ReporteContexto:
        applied = filtros.model_dump(
            mode="json",
            exclude_none=True,
            exclude={"page", "page_size", "sort", "order"},
        )
        return ReporteContexto(
            alcance=filtros.alcance,
            clasificacion=filtros.clasificacion,
            fecha_referencia=None,
            filtros_aplicados=applied,
        )

    def _validate_sort(self, filtros: ReporteEventoFiltros, endpoint: str) -> None:
        if filtros.sort is not None and filtros.sort not in self.SORTS[endpoint]:
            raise ReporteFiltroInconsistenteError(
                f"El campo de orden '{filtros.sort}' no está permitido.", campo="sort"
            )

    async def _validate_event_filters(
        self, id_evento: int, filtros: ReporteEventoFiltros, *, endpoint: str
    ) -> EventoReporteDetalle:
        detail = await self.reportes.get_evento(id_evento)
        if detail is None:
            raise EventoReporteNotFoundError("Evento no encontrado.")
        if endpoint in self.SORTS:
            self._validate_sort(filtros, endpoint)
        if filtros.id_programacion_evento is not None and not await self.reportes.programacion_pertenece(
            id_evento=id_evento,
            id_programacion_evento=filtros.id_programacion_evento,
        ):
            raise ReporteFiltroInconsistenteError(
                "La programación no pertenece al evento.",
                campo="id_programacion_evento",
            )
        if filtros.id_detalle_categoria is not None and not await self.reportes.detalle_categoria_valido(
            id_detalle_categoria=filtros.id_detalle_categoria,
            id_grupo=filtros.id_grupo,
            id_categoria=filtros.id_categoria,
        ):
            raise ReporteFiltroInconsistenteError(
                "La relación grupo-categoría no coincide con el detalle indicado.",
                campo="id_detalle_categoria",
            )
        if (
            filtros.id_detalle_categoria is None
            and filtros.id_grupo is not None
            and filtros.id_categoria is not None
            and not await self.reportes.grupo_categoria_valido(
                id_grupo=filtros.id_grupo, id_categoria=filtros.id_categoria
            )
        ):
            raise ReporteFiltroInconsistenteError(
                "La categoría no pertenece al grupo indicado.", campo="id_categoria"
            )
        if filtros.id_empresa is not None and not await self.reportes.empresa_afiliada(
            id_evento=id_evento, id_empresa=filtros.id_empresa, filtros=filtros
        ):
            raise ReporteFiltroInconsistenteError(
                "La empresa no está afiliada al evento dentro del alcance solicitado.",
                campo="id_empresa",
            )
        return detail

    async def listar_eventos(
        self,
        *,
        search: str | None,
        estado: Any,
        fecha_desde: date | None,
        fecha_hasta: date | None,
        page: int,
        page_size: int,
        sort: str,
        order: ReporteOrden,
    ) -> EventoSelectorResponse:
        if fecha_desde is not None and fecha_hasta is not None and fecha_hasta < fecha_desde:
            raise ReporteFiltroInconsistenteError(
                "fecha_hasta no puede ser anterior a fecha_desde.", campo="fecha_hasta"
            )
        filtros = ReporteEventoFiltros(sort=sort, order=order)
        self._validate_sort(filtros, "eventos")
        rows, total = await self.reportes.list_eventos(
            search=search.strip() if search else None,
            estado=estado,
            fecha_desde=fecha_desde,
            fecha_hasta=fecha_hasta,
            page=page,
            page_size=page_size,
            sort=sort,
            order=order,
        )
        return EventoSelectorResponse(
            items=[
                EventoSelectorItem(
                    id_evento=row["id_evento"],
                    nombre_evento=row["nombre_evento"],
                    descripcion=row["descripcion"],
                    estado=row["estado"],
                    area=AreaReporte(id_area=row["id_area"], nombre_area=row["nombre_area"]),
                    fecha_inicio=row["fecha_inicio"],
                    fecha_fin=row["fecha_fin"],
                )
                for row in rows
            ],
            total=total,
            page=page,
            page_size=page_size,
            pages=self._pages(total, page_size),
        )

    def _build_kpis(self, values: dict[str, int], universe: int) -> ReporteKpis:
        represented = values["total_empresas_representadas"]
        kwargs = {name: self._metric(name, values[name]) for name in self.METRICAS}
        return ReporteKpis(
            **kwargs,
            cobertura_afiliadas=self._coverage(
                "AFILIADAS_EVENTO", represented, values["total_empresas_unicas"]
            ),
            cobertura_universo=self._coverage(
                "ACTIVAS_CLASIFICACION", represented, universe
            ),
        )

    async def obtener_kpis(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> ReporteKpisResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        values = await self.reportes.get_kpi_values(id_evento, filtros)
        universe = await self.reportes.get_universe_count(id_evento, filtros)
        return ReporteKpisResponse(
            contexto=self._context(filtros),
            kpis=self._build_kpis(values, universe),
            disponibilidad_datos=self._availability(),
            advertencias=self._warnings(filtros),
        )

    async def obtener_filtros(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> ReporteFiltrosResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        options = await self.reportes.list_filter_options(id_evento, filtros)
        programs = options["programaciones"]
        classifications = options["clasificaciones"]
        companies = options["empresas"]
        participant_options = options["participante_opciones"]
        cargos = sorted(
            {(row.id_cargo, row.nombre_cargo) for row in participant_options if row.id_cargo is not None},
            key=lambda value: value[1],
        )
        benefits = sorted(
            {(row.id_beneficio, row.nombre_beneficio) for row in participant_options if row.id_beneficio is not None},
            key=lambda value: value[1],
        )
        groups = sorted(
            {(row.id_grupo, row.nombre_grupo) for row in classifications if row.id_grupo is not None},
            key=lambda value: value[1],
        )
        return ReporteFiltrosResponse(
            contexto=self._context(filtros),
            programaciones=[
                FiltroOpcion(
                    id=row.id_programacion_evento,
                    label=f"{row.modalidad.value} - {row.primera_fecha or 'sin fecha'}",
                )
                for row in programs[:100]
            ],
            grupos=[FiltroOpcion(id=item[0], label=item[1]) for item in groups[:100]],
            categorias=[
                CategoriaFiltroOpcion(
                    id_detalle_categoria=row.id_detalle_categoria,
                    id_grupo=row.id_grupo,
                    nombre_grupo=row.nombre_grupo,
                    id_categoria=row.id_categoria,
                    nombre_categoria=row.nombre_categoria,
                )
                for row in classifications[:100]
                if row.id_detalle_categoria is not None
            ],
            empresas=[FiltroOpcion(id=row.id_empresa, label=row.nombre_empresa) for row in companies[:100]],
            cargos=[FiltroOpcion(id=item[0], label=item[1]) for item in cargos[:100]],
            beneficios=[FiltroOpcion(id=item[0], label=item[1]) for item in benefits[:100]],
            modalidades=sorted({row.modalidad for row in programs}, key=lambda item: item.value),
            has_more={
                "programaciones": len(programs) > 100,
                "grupos": len(groups) > 100,
                "categorias": len(classifications) > 100,
                "empresas": len(companies) > 100,
                "cargos": len(cargos) > 100,
                "beneficios": len(benefits) > 100,
            },
        )

    async def _distribution_items(
        self, id_evento: int, filtros: ReporteEventoFiltros, *, by: str
    ) -> list[DistribucionGrupoItem] | list[DistribucionCategoriaItem]:
        rows = await self.reportes.list_distribution(id_evento, filtros, by=by)
        universe = await self.reportes.universe_by_classification(
            id_evento, filtros, by=by
        )
        items: list[Any] = []
        for row in rows:
            key = row["id_grupo"] if by == "grupo" else row["id_detalle_categoria"]
            common = dict(
                afiliaciones=int(row["afiliaciones"]),
                empresas_unicas=int(row["empresas_unicas"]),
                empresas_representadas=int(row["empresas_representadas"]),
                participaciones=int(row["participaciones"]),
                asistencias=int(row["asistencias"]),
                sin_asistencia=int(row["sin_asistencia"]),
                asignaciones_beneficio=int(row["asignaciones_beneficio"]),
                requieren_coordinacion=int(row["requieren_coordinacion"]),
                cobertura_afiliadas=self._coverage(
                    "AFILIADAS_EVENTO", int(row["empresas_representadas"]), int(row["empresas_unicas"])
                ),
                cobertura_universo=self._coverage(
                    "ACTIVAS_CLASIFICACION", int(row["empresas_representadas"]), universe.get(key, 0)
                ),
            )
            if by == "grupo":
                items.append(DistribucionGrupoItem(id_grupo=row["id_grupo"], nombre_grupo=row["nombre_grupo"], **common))
            else:
                items.append(
                    DistribucionCategoriaItem(
                        id_detalle_categoria=row["id_detalle_categoria"],
                        id_grupo=row["id_grupo"],
                        nombre_grupo=row["nombre_grupo"],
                        id_categoria=row["id_categoria"],
                        nombre_categoria=row["nombre_categoria"],
                        **common,
                    )
                )
        return items

    async def obtener_grupos(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> DistribucionGrupoResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        return DistribucionGrupoResponse(
            contexto=self._context(filtros),
            items=await self._distribution_items(id_evento, filtros, by="grupo"),
            advertencias=self._warnings(filtros),
        )

    async def obtener_categorias(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> DistribucionCategoriaResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        return DistribucionCategoriaResponse(
            contexto=self._context(filtros),
            items=await self._distribution_items(id_evento, filtros, by="categoria"),
            advertencias=self._warnings(filtros),
        )

    @staticmethod
    def _classification(row: dict[str, Any], filtros: ReporteEventoFiltros) -> ClasificacionReporte:
        return ClasificacionReporte(
            id_detalle_categoria=row["id_detalle_categoria"],
            id_grupo=row["id_grupo"],
            nombre_grupo=row["nombre_grupo"],
            id_categoria=row["id_categoria"],
            nombre_categoria=row["nombre_categoria"],
            fuente=filtros.clasificacion,
            fecha_referencia=row.get("fecha_referencia") if filtros.clasificacion == ReporteClasificacion.FECHA_PROGRAMACION else None,
        )

    def _company_item(self, row: dict[str, Any], filtros: ReporteEventoFiltros) -> EmpresaReporteItem:
        return EmpresaReporteItem(
            id_empresa=row["id_empresa"], nombre_empresa=row["nombre_empresa"], ruc=row["ruc"],
            razon_social=row["razon_social"], nombre_comercial=row["nombre_comercial"],
            estado_empresa=row["estado_empresa"], estado_afiliacion=row["estado_afiliacion"],
            clasificacion=self._classification(row, filtros), programaciones=list(row["programaciones"] or []),
            afiliaciones=int(row["afiliaciones"]), participantes=int(row["participantes"]),
            asistentes=int(row["asistentes"]), sin_asistencia=int(row["sin_asistencia"]),
            beneficios=int(row["beneficios"]), requieren_coordinacion=int(row["requieren_coordinacion"]),
            qr_generados=int(row["qr_generados"]), qr_enviados=int(row["qr_enviados"]),
            credenciales_impresas=int(row["credenciales_impresas"]),
        )

    async def obtener_empresas(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> EmpresaReporteListResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="empresas")
        rows, total = await self.reportes.list_companies(id_evento, filtros)
        return EmpresaReporteListResponse(
            contexto=self._context(filtros), items=[self._company_item(row, filtros) for row in rows],
            total=total, page=filtros.page, page_size=filtros.page_size,
            pages=self._pages(total, filtros.page_size), advertencias=self._warnings(filtros),
        )

    @staticmethod
    def _attendance_mode(modality: EventoModalidad, attendance: bool) -> str:
        if not attendance:
            return "SIN_ASISTENCIA"
        if modality == EventoModalidad.HIBRIDO:
            return "NO_REGISTRADA"
        return modality.value

    def _participant_item(self, row: dict[str, Any], filtros: ReporteEventoFiltros) -> ParticipanteReporteItem:
        contact = row["id_contacto"] is not None
        names = row["contacto_nombres"] if contact else row["invitado_nombres"]
        surnames = row["contacto_apellidos"] if contact else row["invitado_apellidos"]
        document = row["contacto_numero_documento"] if contact else row["invitado_numero_documento"]
        email = row["contacto_correo"] if contact else row["invitado_correo"]
        phone = row["contacto_celular"] if contact else row["invitado_celular"]
        return ParticipanteReporteItem(
            id_evento_contacto=row["id_evento_contacto"],
            id_programacion_evento=row["id_programacion_evento"],
            modalidad_programacion=row["modalidad_programacion"],
            modalidad_asistencia=self._attendance_mode(row["modalidad_programacion"], row["asistencia"]),
            tipo_participante=ReporteTipoParticipante.CONTACTO if contact else ReporteTipoParticipante.INVITADO,
            id_contacto=row["id_contacto"], id_tipo_documento=row["id_tipo_documento"] if contact else None,
            tipo_documento=row["tipo_documento"] if contact else None, numero_documento=document,
            nombres=names or "", apellidos=surnames or "", nombre_completo=f"{surnames or ''} {names or ''}".strip(),
            genero=row["genero"] if contact else None, id_cargo=row["id_cargo"] if contact else None,
            nombre_cargo=row["nombre_cargo"] if contact else None, correo=email, celular=phone,
            id_empresa=row["id_empresa"], nombre_empresa=row["nombre_empresa"],
            clasificacion=self._classification(row, filtros), estado_participacion=row["estado_participacion"],
            asistencia=row["asistencia"], hora_ingreso=row["hora_ingreso"], id_beneficio=row["id_beneficio"],
            nombre_beneficio=row["nombre_beneficio"], tipo_calculo_beneficio=row["tipo_calculo_beneficio"],
            requiere_coordinacion=row["requiere_coordinacion"], qr_generado=row["id_participante_qr"] is not None,
            qr_enviado=row["qr_fecha_envio"] is not None, qr_estado=row["qr_estado"],
            credencial_impresa=row["credencial_impresa"],
        )

    async def obtener_participantes(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> ParticipanteReporteListResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        rows, total = await self.reportes.list_participations(id_evento, filtros)
        return ParticipanteReporteListResponse(
            contexto=self._context(filtros), items=[self._participant_item(row, filtros) for row in rows],
            total=total, page=filtros.page, page_size=filtros.page_size,
            pages=self._pages(total, filtros.page_size), advertencias=self._warnings(filtros),
        )

    async def obtener_detalle(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> DetalleReporteListResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="detalle")
        rows, total = await self.reportes.list_participations(id_evento, filtros)
        items = []
        for row in rows:
            participant = self._participant_item(row, filtros).model_dump()
            items.append(
                DetalleReporteItem(
                    **participant,
                    id_evento=row["id_evento"], nombre_evento=row["nombre_evento"],
                    estado_evento=row["estado_evento"], estado_programacion=row["estado_programacion"],
                    primera_fecha_programacion=row["primera_fecha_programacion"],
                )
            )
        return DetalleReporteListResponse(
            contexto=self._context(filtros), items=items, total=total, page=filtros.page,
            page_size=filtros.page_size, pages=self._pages(total, filtros.page_size),
            advertencias=self._warnings(filtros),
        )

    async def obtener_programaciones(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> ProgramacionReporteResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        rows, days, responsibles = await self.reportes.list_programs(id_evento, filtros)
        days_by_program: dict[int, list[Any]] = defaultdict(list)
        for day in days:
            days_by_program[day.id_programacion_evento].append(day)
        responsible_by_program: dict[int, list[Any]] = defaultdict(list)
        for responsible in responsibles:
            responsible_by_program[responsible.id_programacion_evento].append(responsible)
        items = []
        for row in rows:
            program_days = days_by_program[row["id_programacion_evento"]]
            items.append(
                ProgramacionReporteItem(
                    id_programacion_evento=row["id_programacion_evento"], modalidad=row["modalidad"],
                    estado=row["estado"], enlace_general=row["enlace_general"],
                    lugar=LugarReporte(
                        id_lugar=row["id_lugar"], pais=row["pais"], provincia=row["provincia"],
                        distrito=row["distrito"], direccion=row["direccion"], estado=row["lugar_estado"],
                    ) if row["id_lugar"] is not None else None,
                    primera_fecha=min((day.fecha for day in program_days), default=None),
                    ultima_fecha=max((day.fecha for day in program_days), default=None),
                    cantidad_dias=len(program_days),
                    dias=[DiaProgramacionReporte.model_validate(day, from_attributes=True) for day in program_days],
                    responsables=[
                        ResponsableProgramacionReporte(
                            id_responsable_evento=item.id_responsable_evento, id_usuario=item.id_usuario,
                            nombre_usuario=item.nombre_usuario, nombres=item.nombres,
                            apellidos=item.apellidos, estado=item.estado,
                        ) for item in responsible_by_program[row["id_programacion_evento"]]
                    ],
                    empresas_afiliadas=int(row["empresas_afiliadas"]),
                    participaciones=int(row["participaciones"]), asistencias=int(row["asistencias"]),
                )
            )
        return ProgramacionReporteResponse(contexto=self._context(filtros), items=items)

    async def obtener_series(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> SerieMensualResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        rows = await self.reportes.list_series(id_evento, filtros)
        by_metric: dict[str, list[SerieDato]] = defaultdict(list)
        for metric, period, value in rows:
            by_metric[metric].append(SerieDato(periodo=period, valor=value))
        series = [
            SerieReporte(metrica=name, fuente_temporal=source, datos=by_metric.get(name, []))
            for name, source in self.SERIES_FUENTE.items()
        ]
        return SerieMensualResponse(
            contexto=self._context(filtros), series=series, advertencias=self._warnings(filtros)
        )

    async def obtener_beneficios(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> BeneficioReporteResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        details = await self.reportes.list_policy_details(id_evento)
        usage, coordination = await self.reportes.benefit_usage(id_evento, filtros)
        per_program, unique_companies = await self.reportes.policy_company_counts(id_evento, filtros)
        items = []
        for detail, benefit, category in details:
            key = (benefit.id_beneficio, category.id_categoria)
            used = usage.get(key, {"asignaciones": 0, "usados": 0})
            applicable = (
                unique_companies.get(category.id_categoria, 0)
                if benefit.tipo_calculo == TipoCalculoBeneficio.POR_ANIO
                else per_program.get(category.id_categoria, 0)
            )
            available = benefit.tipo_calculo != TipoCalculoBeneficio.SIN_BENEFICIO
            capacity = detail.entradas_gratuitas * benefit.personas_por_asignacion * applicable if available else None
            if benefit.tipo_calculo == TipoCalculoBeneficio.POR_EVENTO:
                consumed = used["asignaciones"]
            elif benefit.tipo_calculo == TipoCalculoBeneficio.POR_ANIO:
                consumed = used["usados"]
            else:
                consumed = None
            items.append(
                BeneficioReporteItem(
                    id_detalle_politica_evento=detail.id_detalle_politica_evento,
                    id_beneficio=benefit.id_beneficio, nombre=benefit.nombre,
                    tipo_calculo=benefit.tipo_calculo, id_categoria=category.id_categoria,
                    nombre_categoria=category.nombre_categoria,
                    entradas_gratuitas=detail.entradas_gratuitas,
                    personas_por_asignacion=benefit.personas_por_asignacion,
                    empresas_aplicables=applicable, cupos_teoricos=capacity,
                    cupos_utilizados=consumed,
                    cupos_disponibles=max(capacity - consumed, 0) if capacity is not None and consumed is not None else None,
                    asignaciones=used["asignaciones"],
                    requieren_coordinacion=coordination.get(key, 0), disponible=available,
                )
            )
        return BeneficioReporteResponse(
            contexto=self._context(filtros), items=items, advertencias=self._warnings(filtros)
        )

    async def obtener_acreditacion(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> AcreditacionReporteResponse:
        await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        values, reprints = await self.reportes.accreditation(id_evento, filtros)
        return AcreditacionReporteResponse(
            contexto=self._context(filtros),
            qr_generados=self._metric("qr_generados", values["qr_generados"]),
            qr_enviados=self._metric("qr_enviados", values["qr_enviados"]),
            qr_pendientes=self._metric("qr_pendientes", values["qr_pendientes"]),
            credenciales_impresas=self._metric("credenciales_impresas", values["credenciales_impresas"]),
            credenciales_no_impresas=self._metric("credenciales_no_impresas", values["credenciales_no_impresas"]),
            reimpresiones=MetricaValor(
                valor=reprints, calidad=CalidadDato.INFERIDO, grano="auditoría",
                definicion="Auditorías REIMPRESION_CREDENCIAL relacionadas con participaciones filtradas.",
            ),
        )

    @staticmethod
    def _event_header(detail: EventoReporteDetalle) -> EventoReporteCabecera:
        return EventoReporteCabecera(
            id_evento=detail.evento.id_evento, nombre_evento=detail.evento.nombre_evento,
            descripcion=detail.evento.descripcion, estado=detail.evento.estado,
            area=AreaReporte(id_area=detail.area.id_area, nombre_area=detail.area.nombre_area),
            fecha_inicio=detail.politica.fecha_inicio, fecha_fin=detail.politica.fecha_fin,
            flyer_url=detail.evento.flyer_url,
        )

    async def obtener_dashboard(
        self, id_evento: int, filtros: ReporteEventoFiltros
    ) -> DashboardReporteResponse:
        detail = await self._validate_event_filters(id_evento, filtros, endpoint="participantes")
        values = await self.reportes.get_kpi_values(id_evento, filtros)
        universe = await self.reportes.get_universe_count(id_evento, filtros)
        series_rows = await self.reportes.list_series(id_evento, filtros)
        principal = [SerieDato(periodo=period, valor=value) for metric, period, value in series_rows if metric == "participaciones_programadas"]
        groups = await self._distribution_items(id_evento, filtros, by="grupo")
        categories = await self._distribution_items(id_evento, filtros, by="categoria")
        return DashboardReporteResponse(
            evento=self._event_header(detail), contexto=self._context(filtros),
            kpis=self._build_kpis(values, universe),
            serie_principal=SerieReporte(
                metrica="participaciones_programadas", fuente_temporal="PRIMER_DIA_PROGRAMACION", datos=principal
            ),
            distribuciones={
                "grupos": [item.model_dump(mode="json") for item in groups],
                "categorias": [item.model_dump(mode="json") for item in categories],
            },
            disponibilidad_datos=self._availability(), advertencias=self._warnings(filtros),
        )

    async def _can_export_pii(self, actor: Usuario) -> bool:
        return await self.usuarios.has_permission(
            id_rol=actor.id_rol, modulo=MODULO_REPORTES, permiso=PERMISO_PII
        )

    @staticmethod
    def _write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 50)
            sheet.column_dimensions[column[0].column_letter].width = width

    async def exportar(
        self, id_evento: int, filtros: ReporteEventoFiltros, actor: Usuario
    ) -> bytes:
        detail = await self._validate_event_filters(id_evento, filtros, endpoint="detalle")
        participant_rows, total_participants = await self.reportes.list_participations(
            id_evento, filtros, paginated=False
        )
        if total_participants > settings.report_export_sync_max_rows:
            raise ReporteExportLimitError(
                f"La exportación contiene {total_participants} filas y supera el límite síncrono de {settings.report_export_sync_max_rows}."
            )
        company_rows, _ = await self.reportes.list_companies(id_evento, filtros, paginated=False)
        values = await self.reportes.get_kpi_values(id_evento, filtros)
        universe = await self.reportes.get_universe_count(id_evento, filtros)
        programs = await self.obtener_programaciones(id_evento, filtros)
        agenda = await self.reportes.list_agenda(id_evento, filtros)
        benefits = await self.obtener_beneficios(id_evento, filtros)
        include_pii = await self._can_export_pii(actor)

        workbook = Workbook()
        workbook.remove(workbook.active)
        kpis = self._build_kpis(values, universe)
        self._write_sheet(
            workbook, "Resumen", ["Métrica", "Valor"],
            [[name, getattr(kpis, name).valor] for name in self.METRICAS]
            + [["cobertura_afiliadas", kpis.cobertura_afiliadas.porcentaje], ["cobertura_universo", kpis.cobertura_universo.porcentaje]],
        )
        self._write_sheet(
            workbook, "Programaciones",
            ["id_programacion", "modalidad", "estado", "primera_fecha", "ultima_fecha", "dias", "empresas", "participaciones", "asistencias"],
            [[item.id_programacion_evento, item.modalidad.value, item.estado.value, item.primera_fecha, item.ultima_fecha, item.cantidad_dias, item.empresas_afiliadas, item.participaciones, item.asistencias] for item in programs.items],
        )
        self._write_sheet(
            workbook, "Agenda", ["id_dia", "id_programacion", "fecha", "hora_inicio", "hora_fin", "estado"],
            [[row["id_detalle_programacion"], row["id_programacion_evento"], row["fecha"], row["hora_inicio"], row["hora_fin"], row["estado"]] for row in agenda],
        )
        self._write_sheet(
            workbook, "Empresas", ["id_empresa", "ruc", "empresa", "grupo", "categoria", "afiliaciones", "participantes", "asistentes"],
            [[row["id_empresa"], row["ruc"], row["nombre_empresa"], row["nombre_grupo"], row["nombre_categoria"], row["afiliaciones"], row["participantes"], row["asistentes"]] for row in company_rows],
        )
        participant_headers = ["id_evento_contacto", "id_programacion", "tipo", "nombres", "apellidos", "empresa", "asistencia", "beneficio", "qr_generado", "qr_enviado", "credencial_impresa"]
        if include_pii:
            participant_headers[5:5] = ["documento", "correo", "celular"]
        participant_export = []
        for row in participant_rows:
            item = self._participant_item(row, filtros)
            values_row: list[Any] = [item.id_evento_contacto, item.id_programacion_evento, item.tipo_participante.value, item.nombres, item.apellidos]
            if include_pii:
                values_row.extend([item.numero_documento, item.correo, item.celular])
            values_row.extend([item.nombre_empresa, item.asistencia, item.nombre_beneficio, item.qr_generado, item.qr_enviado, item.credencial_impresa])
            participant_export.append(values_row)
        self._write_sheet(workbook, "Participantes", participant_headers, participant_export)
        self._write_sheet(
            workbook, "Beneficios", ["id_beneficio", "beneficio", "tipo_calculo", "categoria", "empresas_aplicables", "cupos_teoricos", "cupos_utilizados", "cupos_disponibles", "asignaciones"],
            [[item.id_beneficio, item.nombre, item.tipo_calculo.value, item.nombre_categoria, item.empresas_aplicables, item.cupos_teoricos, item.cupos_utilizados, item.cupos_disponibles, item.asignaciones] for item in benefits.items],
        )
        module = await self.usuarios.get_module_by_name(MODULO_REPORTES)
        try:
            await self.auditoria.create(
                id_usuario=actor.id_usuario,
                id_modulo=module.id_modulo if module else None,
                entidad="evento",
                id_entidad=id_evento,
                accion="EXPORTAR_REPORTE_EVENTO",
                valor_nuevo={
                    "formato": "XLSX", "filas_participantes": total_participants,
                    "incluye_datos_personales": include_pii,
                    "filtros": self._context(filtros).filtros_aplicados,
                    "evento": detail.evento.nombre_evento,
                },
            )
            await self.db.commit()
        except Exception:
            await self.db.rollback()
            raise
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()
