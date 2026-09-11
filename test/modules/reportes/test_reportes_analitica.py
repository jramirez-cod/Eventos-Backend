from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.config import settings
from app.modules.auditoria.models import Auditoria
from test.modules.reportes.conftest import create_report_actor, seed_report_context


pytestmark = pytest.mark.asyncio


def _recursive_keys(value: object) -> set[str]:
    if isinstance(value, dict):
        result = {str(key).lower() for key in value}
        for nested in value.values():
            result.update(_recursive_keys(nested))
        return result
    if isinstance(value, list):
        result: set[str] = set()
        for nested in value:
            result.update(_recursive_keys(nested))
        return result
    return set()


async def test_series_beneficios_y_acreditacion(client, session_factory) -> None:
    async with session_factory() as session:
        context = await seed_report_context(session)
        event_id = context["event"].id_evento
    headers = context["headers"]
    series = await client.get(f"/api/v1/reportes/eventos/{event_id}/series/mensual", headers=headers)
    benefits = await client.get(f"/api/v1/reportes/eventos/{event_id}/beneficios", headers=headers)
    accreditation = await client.get(f"/api/v1/reportes/eventos/{event_id}/acreditacion", headers=headers)
    assert series.status_code == benefits.status_code == accreditation.status_code == 200
    by_name = {item["metrica"]: item for item in series.json()["series"]}
    assert by_name["participaciones_programadas"]["fuente_temporal"] == "PRIMER_DIA_PROGRAMACION"
    assert by_name["participaciones_programadas"]["datos"] == [
        {"periodo": "2026-01", "valor": 1},
        {"periodo": "2026-02", "valor": 1},
    ]
    benefit_types = {item["tipo_calculo"]: item for item in benefits.json()["items"]}
    assert benefit_types["POR_EVENTO"]["asignaciones"] == 1
    assert benefit_types["POR_EVENTO"]["cupos_utilizados"] == 1
    assert benefit_types["POR_ANIO"]["cupos_utilizados"] == 1
    assert benefit_types["SIN_BENEFICIO"]["disponible"] is False
    assert benefit_types["SIN_BENEFICIO"]["cupos_teoricos"] is None
    data = accreditation.json()
    assert data["qr_generados"]["valor"] == 2
    assert data["qr_enviados"]["valor"] == 1
    assert data["qr_pendientes"]["valor"] == 1
    assert data["reimpresiones"]["valor"] == 1
    assert data["reimpresiones"]["calidad"] == "INFERIDO"


async def test_consistencia_dashboard_kpi_detalle_excel(client, session_factory) -> None:
    async with session_factory() as session:
        context = await seed_report_context(session)
        event_id = context["event"].id_evento
    headers = context["headers"]
    queries = (
        "",
        f"id_grupo={context['group_a'].id_grupo}&id_categoria={context['category_a'].id_categoria}",
        f"id_empresa={context['company'].id_empresa}",
        f"id_programacion_evento={context['programs'][0].id_programacion_evento}",
        "asistencia=true",
        "alcance=HISTORICO",
    )
    for query in queries:
        prefix = f"?{query}&" if query else "?"
        dashboard = await client.get(f"/api/v1/reportes/eventos/{event_id}/dashboard{prefix}page_size=100", headers=headers)
        kpis = await client.get(f"/api/v1/reportes/eventos/{event_id}/kpis{prefix}page_size=100", headers=headers)
        detail = await client.get(f"/api/v1/reportes/eventos/{event_id}/detalle{prefix}page_size=100", headers=headers)
        export = await client.get(f"/api/v1/reportes/eventos/{event_id}/exportar{prefix}page_size=100", headers=headers)
        assert all(response.status_code == 200 for response in (dashboard, kpis, detail, export)), query
        expected = kpis.json()["kpis"]["total_participaciones"]["valor"]
        assert dashboard.json()["kpis"]["total_participaciones"]["valor"] == expected
        assert len(detail.json()["items"]) == expected
        workbook = load_workbook(BytesIO(export.content), read_only=True, data_only=True)
        assert workbook["Participantes"].max_row - 1 == expected
        summary = {row[0]: row[1] for row in workbook["Resumen"].iter_rows(min_row=2, values_only=True)}
        assert summary["total_participaciones"] == expected


async def test_exportacion_sin_permiso_pii_oculta_columnas(client, session_factory) -> None:
    async with session_factory() as session:
        context = await seed_report_context(session)
        _, headers = await create_report_actor(
            session,
            permissions=("CONSULTAR_REPORTE_EVENTO", "EXPORTAR_REPORTE_EVENTO"),
        )
        event_id = context["event"].id_evento
    response = await client.get(f"/api/v1/reportes/eventos/{event_id}/exportar", headers=headers)
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    headers_row = [cell.value for cell in next(workbook["Participantes"].iter_rows(max_row=1))]
    assert "documento" not in headers_row
    assert "correo" not in headers_row
    assert "celular" not in headers_row


async def test_json_y_excel_no_exponen_secretos(client, session_factory) -> None:
    async with session_factory() as session:
        context = await seed_report_context(session)
        event_id = context["event"].id_evento
    forbidden = {
        "codigo_seguro", "codigo_hash", "password", "password_hash", "token",
        "access_token", "refresh_token", "smtp_app_password",
    }
    for suffix in ("dashboard", "participantes", "detalle", "acreditacion"):
        response = await client.get(f"/api/v1/reportes/eventos/{event_id}/{suffix}", headers=context["headers"])
        assert response.status_code == 200
        assert not (_recursive_keys(response.json()) & forbidden)
        assert "qr-secret" not in response.text.lower()
    export = await client.get(f"/api/v1/reportes/eventos/{event_id}/exportar", headers=context["headers"])
    workbook = load_workbook(BytesIO(export.content), read_only=True, data_only=True)
    text = " ".join(str(cell) for sheet in workbook for row in sheet.iter_rows(values_only=True) for cell in row if cell is not None).lower()
    assert not any(secret in text for secret in forbidden | {"qr-secret"})


async def test_limite_de_exportacion_es_configurable(client, session_factory, monkeypatch) -> None:
    async with session_factory() as session:
        context = await seed_report_context(session)
        event_id = context["event"].id_evento
    monkeypatch.setattr(settings, "report_export_sync_max_rows", 1)
    response = await client.get(f"/api/v1/reportes/eventos/{event_id}/exportar", headers=context["headers"])
    assert response.status_code == 413
    assert response.json()["detail"]["codigo"] == "REPORTE_EXPORTACION_EXCEDE_LIMITE"
    async with session_factory() as session:
        audits = await session.scalar(
            select(func.count()).select_from(Auditoria).where(
                Auditoria.accion == "EXPORTAR_REPORTE_EVENTO",
                Auditoria.id_entidad == str(event_id),
            )
        )
        assert audits == 0
