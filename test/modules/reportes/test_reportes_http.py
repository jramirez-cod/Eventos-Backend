from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.modules.auditoria.models import Auditoria
from test.modules.reportes.conftest import create_report_actor, seed_report_context


pytestmark = pytest.mark.asyncio


async def test_openapi_expone_todos_los_endpoints(client) -> None:
    response = await client.get("/openapi.json")
    paths = response.json()["paths"]
    expected = {
        "/api/v1/reportes/eventos",
        "/api/v1/reportes/eventos/{id_evento}/filtros",
        "/api/v1/reportes/eventos/{id_evento}/kpis",
        "/api/v1/reportes/eventos/{id_evento}/grupos",
        "/api/v1/reportes/eventos/{id_evento}/categorias",
        "/api/v1/reportes/eventos/{id_evento}/empresas",
        "/api/v1/reportes/eventos/{id_evento}/participantes",
        "/api/v1/reportes/eventos/{id_evento}/programaciones",
        "/api/v1/reportes/eventos/{id_evento}/detalle",
        "/api/v1/reportes/eventos/{id_evento}/dashboard",
        "/api/v1/reportes/eventos/{id_evento}/series/mensual",
        "/api/v1/reportes/eventos/{id_evento}/beneficios",
        "/api/v1/reportes/eventos/{id_evento}/acreditacion",
        "/api/v1/reportes/eventos/{id_evento}/exportar",
    }
    assert expected <= set(paths)
    assert "/api/v1/reportes/eventos/{id_evento}/completo" not in paths


async def test_endpoints_exigen_autenticacion(client) -> None:
    response = await client.get("/api/v1/reportes/eventos")
    assert response.status_code == 401


async def test_rbac_admin_personal_sin_permiso_e_inactivo(client, session_factory) -> None:
    async with session_factory() as session:
        context = await seed_report_context(session)
        personal, personal_headers = await create_report_actor(
            session, permissions=("CONSULTAR_REPORTE_EVENTO",)
        )
        _, no_permission_headers = await create_report_actor(session, permissions=())
        _, inactive_headers = await create_report_actor(
            session, permissions=("CONSULTAR_REPORTE_EVENTO",), active=False
        )
        event_id = context["event"].id_evento

    assert (await client.get(f"/api/v1/reportes/eventos/{event_id}/dashboard", headers=context["headers"])).status_code == 200
    assert (await client.get(f"/api/v1/reportes/eventos/{event_id}/participantes", headers=context["headers"])).status_code == 200
    assert (await client.get(f"/api/v1/reportes/eventos/{event_id}/exportar", headers=context["headers"])).status_code == 200
    assert (await client.get(f"/api/v1/reportes/eventos/{event_id}/dashboard", headers=personal_headers)).status_code == 200
    assert (await client.get(f"/api/v1/reportes/eventos/{event_id}/participantes", headers=personal_headers)).status_code == 403
    assert (await client.get(f"/api/v1/reportes/eventos/{event_id}/exportar", headers=personal_headers)).status_code == 403
    assert (await client.get(f"/api/v1/reportes/eventos/{event_id}/dashboard", headers=no_permission_headers)).status_code == 403
    assert (await client.get(f"/api/v1/reportes/eventos/{event_id}/dashboard", headers=inactive_headers)).status_code == 401
    assert personal.estado is True


async def test_errores_404_400_y_422(client, session_factory) -> None:
    async with session_factory() as session:
        context = await seed_report_context(session)
        event_id = context["event"].id_evento
        headers = context["headers"]

    not_found = await client.get("/api/v1/reportes/eventos/999999/kpis", headers=headers)
    assert not_found.status_code == 404
    inconsistent = await client.get(
        f"/api/v1/reportes/eventos/{event_id}/kpis?id_programacion_evento=999999",
        headers=headers,
    )
    assert inconsistent.status_code == 400
    assert inconsistent.json()["detail"]["campo"] == "id_programacion_evento"
    invalid = await client.get(
        f"/api/v1/reportes/eventos/{event_id}/empresas?page_size=101",
        headers=headers,
    )
    assert invalid.status_code == 422


async def test_todos_los_endpoints_responden_y_exportacion_audita(client, session_factory) -> None:
    async with session_factory() as session:
        context = await seed_report_context(session)
        event_id = context["event"].id_evento
        headers = context["headers"]
    suffixes = (
        "filtros", "kpis", "grupos", "categorias", "empresas", "participantes",
        "programaciones", "detalle", "dashboard", "series/mensual", "beneficios", "acreditacion",
    )
    selector = await client.get("/api/v1/reportes/eventos", headers=headers)
    assert selector.status_code == 200
    for suffix in suffixes:
        response = await client.get(f"/api/v1/reportes/eventos/{event_id}/{suffix}", headers=headers)
        assert response.status_code == 200, (suffix, response.text)
    export = await client.get(f"/api/v1/reportes/eventos/{event_id}/exportar", headers=headers)
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content), read_only=True)
    assert workbook.sheetnames == ["Resumen", "Programaciones", "Agenda", "Empresas", "Participantes", "Beneficios"]
    async with session_factory() as session:
        count = await session.scalar(
            select(func.count()).select_from(Auditoria).where(
                Auditoria.accion == "EXPORTAR_REPORTE_EVENTO",
                Auditoria.id_entidad == str(event_id),
            )
        )
        assert count == 1
